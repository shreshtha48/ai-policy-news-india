"""Write CSVs (the dashboard contract) + one XLSX workbook for humans."""
from __future__ import annotations

import csv
import json
from pathlib import Path

EVENT_COLS = ["event_id", "state_code", "state_name", "city", "category", "sector", "title", "summary",
              "event_date", "date_precision", "first_reported_date", "primary_source_name",
              "primary_source_url", "source_count", "actors", "amount_inr_crore"]
ARTICLE_COLS = ["article_id", "event_id", "state_code", "title", "source_name", "source_url",
                "published_date", "fetched_via", "is_primary", "state_basis", "matched_terms"]
REJECT_COLS = ["reason", "state_code", "title", "source_name", "source_url", "published_date", "fetched_via"]


def write_csv(path: Path, rows: list[dict], cols: list[str]):
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=cols, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)


def write_xlsx(path: Path, sheets: dict[str, tuple[list[str], list[dict]]]):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    wb.remove(wb.active)
    for name, (cols, rows) in sheets.items():
        ws = wb.create_sheet(name)
        ws.append(cols)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in rows:
            ws.append([r.get(c, "") for c in cols])
        ws.freeze_panes = "A2"
        for i, c in enumerate(cols, 1):
            width = min(60, max(len(c), *(len(str(r.get(c, ""))) for r in rows[:200])) + 2) if rows else len(c) + 2
            ws.column_dimensions[ws.cell(1, i).column_letter].width = width
    wb.save(path)


def export(out_dir: Path, events, articles, rejected, reference_dir: Path, summary: dict):
    write_csv(out_dir / "events.csv", events, EVENT_COLS)
    write_csv(out_dir / "articles.csv", articles, ARTICLE_COLS)
    write_csv(out_dir / "rejected.csv", rejected, REJECT_COLS)
    (out_dir / "run_summary.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")
    sheets = {"events": (EVENT_COLS, events), "articles": (ARTICLE_COLS, articles),
              "rejected": (REJECT_COLS, rejected)}
    for ref in ("states", "categories"):
        p = reference_dir / f"{ref}.csv"
        if p.exists():
            with open(p, encoding="utf-8") as f:
                rows = list(csv.DictReader(f))
            sheets[ref] = (list(rows[0].keys()) if rows else [], rows)
    write_xlsx(out_dir / "ai_policy_india_states.xlsx", sheets)
